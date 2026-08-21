#!/usr/bin/env python3
"""Build a costed estimate workbook from a Change Impact Assessment markdown.

The markdown written by the /ia skill is the source of truth. This script reads
its "## Affected Artifacts" table and prices each row, so the workbook can never
disagree with the assessment the user was shown.

Every rate lives on an Assumptions sheet and the Detail and Summary sheets
reference it by cell, so the workbook stays live: change a rate in Excel and all
totals recalculate without re-running this script. That matters in front of a
client — the conversation moves to the rates, which is negotiable, instead of to
the total, which looks arbitrary.

Required input table (see references/change-impact-analysis.md §7):

    ## Affected Artifacts
    | Library | Object | Type | Attribute | Impact | Lines | Notes |
    |---|---|---|---|---|---|---|
    | CASELIB | ORDENT | *PGM | RPGLE | CHANGE | 1840 | rewrites CUSTNO |

    Impact is one of CHANGE / RECOMPILE / REBUILD / REVIEW / NONE.

Requires:  pip install openpyxl
Usage:     python build_change_estimate.py ASSESSMENT.md [--xlsx] [--narrowing]
                                           [--outdir DIR]
"""
from __future__ import annotations

import argparse
import datetime as _dt
import re
import sys
from pathlib import Path

# Section numbering is optional — "## Affected Artifacts" and "## 3. Affected
# Artifacts" must both match, since the report template numbers its sections.
HEADING = re.compile(r"^##\s+(?:\d+(?:\.\d+)*\.?\s+)?affected\s+artifacts\b", re.I)
WANTED = ["library", "object", "type", "attribute", "impact", "lines", "notes"]

# Assumptions sheet layout — every formula below addresses these by row, so the
# two must move together.
A_PGM_S, A_PGM_M, A_PGM_L = 2, 3, 4
A_PGM_RECOMPILE = 5
A_PF, A_LF, A_DSPF, A_PRTF, A_COPYBOOK, A_OTHER, A_REVIEW = 6, 7, 8, 9, 10, 11, 12
A_DESIGN, A_CUTOVER, A_UNIT, A_SYSTEST, A_UAT, A_DATACONV, A_CONTINGENCY = (
    14, 15, 16, 17, 18, 19, 20
)
A_BAND_S, A_BAND_M, A_HOURS_DAY = 22, 23, 24

ASSUMPTIONS = [
    ("Program change — S (small)", 1.5,
     "hours; line count at or below the S threshold; <= 500 source lines"),
    ("Program change — M (medium)", 3.0, "hours; 501-2,000 lines"),
    ("Program change — L (large)", 5.0, "hours; > 2,000 lines"),
    ("Program — recompile only", 0.15, "hours; level-check rebuild, no code edit"),
    ("Physical file / DDS-DDL change", 0.15, "hours"),
    ("Logical file / index rebuild", 0.15, "hours each"),
    ("Display file (DSPF) change", 1.0, "hours"),
    ("Printer file (PRTF) change", 1.0, "hours"),
    ("Copybook change", 0.15, "hours"),
    ("Other object change", 0.15, "hours"),
    ("Review-only item", 0.5, "hours; inspect and confirm no change needed"),
    ("", None, ""),
    ("Impact review & change design", 2.0, "hours, fixed"),
    ("Cutover / implementation", 2.0, "hours, fixed"),
    ("Unit test per changed program", 0.5, "hours"),
    ("System / regression test", 0.15, "share of development"),
    ("UAT support", 0.10, "share of development"),
    ("Data conversion", 2.0, "hours; counted only when the change can lose data"),
    ("Contingency", 0.12, "share of subtotal; raise deliberately, not by default"),
    ("", None, ""),
    ("Band threshold — S up to", 500, "source lines"),
    ("Band threshold — M up to", 2000, "source lines"),
    ("Hours per working day", 8, ""),
]

# Basis keys are written into the Detail sheet and counted by Summary formulas,
# so they must stay literal and stable.
RATE_ROW = {
    "PGM_CHANGE": None,  # banded — resolved by formula against the Band column
    "PGM_RECOMPILE": A_PGM_RECOMPILE,
    "PF_CHANGE": A_PF,
    "LF_REBUILD": A_LF,
    "DSPF_CHANGE": A_DSPF,
    "PRTF_CHANGE": A_PRTF,
    "COPYBOOK_CHANGE": A_COPYBOOK,
    "OTHER_CHANGE": A_OTHER,
    "REVIEW_ONLY": A_REVIEW,
    "NO_COST": None,
}

PROGRAM_TYPES = {"*PGM", "*SRVPGM", "*MODULE"}
PROGRAM_ATTRS = {"RPGLE", "SQLRPGLE", "RPG", "SQLRPG", "RPG38", "RPGIII",
                 "CLLE", "CLP", "CL", "CBLLE", "CBL", "PGM", "SRVPGM", "MODULE"}
COPYBOOK_ATTRS = {"RPGLEINC", "CPYBK", "COPYBOOK", "INC", "SQLINC"}
LF_ATTRS = {"LF", "INDEX", "VIEW", "LF38", "DDS_LF", "MQT"}
PF_ATTRS = {"PF", "PF-DATA", "PF-SRC", "PF38", "TABLE", "PHYSICAL"}


def split_row(line: str) -> list[str]:
    """Split a markdown table row, honouring escaped pipes."""
    line = line.strip().replace(r"\|", "\x00")
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.replace("\x00", "|").strip() for c in line.split("|")]


def is_divider(line: str) -> bool:
    cells = split_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "")


def strip_md(text: str) -> str:
    """Flatten inline markdown — cells land in Excel as plain text."""
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = text.replace("`", "")
    # Object types (*PGM) start with an asterisk the emphasis stripper would eat.
    text = re.sub(r"\*([A-Z][A-Z0-9]{1,9})\b", "\x01\\1", text)
    text = re.sub(r"\*\*|\*", "", text)
    return text.replace("\x01", "*").strip()


def parse(md_path: Path) -> tuple[str, list[dict]]:
    """Return (title, artifact rows) from the Affected Artifacts table."""
    lines = md_path.read_text(encoding="utf-8").splitlines()
    title = md_path.stem.replace("_", " ")
    rows: list[dict] = []
    in_section = False
    i = 0

    while i < len(lines):
        line = lines[i].strip()

        if line.startswith("# ") and not line.startswith("##"):
            title = strip_md(line[2:])
        elif line.startswith("## "):
            in_section = bool(HEADING.match(line))
        elif in_section and line.startswith("|") and i + 1 < len(lines) \
                and is_divider(lines[i + 1]):
            header = [strip_md(c).lower() for c in split_row(line)]
            idx = {name: header.index(name) for name in WANTED if name in header}
            missing = [n for n in WANTED[:5] if n not in idx]
            if missing:
                sys.exit(f"error: Affected Artifacts table is missing column(s): "
                         f"{', '.join(missing)}")
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_divider(lines[i]):
                    cells = [strip_md(c) for c in split_row(lines[i])]
                    rows.append({
                        name: (cells[pos] if pos < len(cells) else "")
                        for name, pos in idx.items()
                    })
                i += 1
            in_section = False
            continue
        i += 1

    return title, rows


def classify(row: dict) -> str:
    """Map an artifact row onto a rate key."""
    impact = row.get("impact", "").upper()
    otype = row.get("type", "").upper()
    attr = row.get("attribute", "").upper()

    if impact in ("NONE", ""):
        return "NO_COST"
    if impact == "REVIEW":
        return "REVIEW_ONLY"
    # A level-check rebuild costs the same whatever the artifact is — without
    # this, a DSPF or PF marked RECOMPILE would be priced at the full change
    # rate. REBUILD keeps falling through to the type dispatch so that logical
    # files still land on LF_REBUILD.
    if impact == "RECOMPILE":
        return "PGM_RECOMPILE"

    is_program = otype in PROGRAM_TYPES or attr in PROGRAM_ATTRS
    if attr in COPYBOOK_ATTRS:
        return "COPYBOOK_CHANGE"
    if is_program:
        return "PGM_CHANGE" if impact == "CHANGE" else "PGM_RECOMPILE"
    if attr in LF_ATTRS:
        return "LF_REBUILD"
    if attr in PF_ATTRS:
        return "PF_CHANGE"
    if attr == "DSPF":
        return "DSPF_CHANGE"
    if attr == "PRTF":
        return "PRTF_CHANGE"
    if otype == "*FILE":
        return "PF_CHANGE"
    return "OTHER_CHANGE"


def to_int(text: str):
    digits = re.sub(r"[^0-9]", "", text or "")
    return int(digits) if digits else None


def build(out: Path, title: str, rows: list[dict], narrowing: bool) -> dict:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("openpyxl is required:  pip install openpyxl")

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    dim = Font(italic=True, color="44546A")
    money_fill = PatternFill("solid", fgColor="FFF2CC")

    # ---------------------------------------------------------------- Detail
    det = wb.create_sheet("Estimate Detail")
    headers = ["Library", "Object", "Type", "Attribute", "Impact",
               "Lines", "Band", "Basis", "Hours", "Notes"]
    det.append(headers)
    for c in range(1, len(headers) + 1):
        det.cell(row=1, column=c).font = head_font
        det.cell(row=1, column=c).fill = head_fill

    counts: dict[str, int] = {}
    for n, row in enumerate(rows, start=2):
        basis = classify(row)
        counts[basis] = counts.get(basis, 0) + 1
        det.cell(row=n, column=1, value=row.get("library", ""))
        det.cell(row=n, column=2, value=row.get("object", ""))
        det.cell(row=n, column=3, value=row.get("type", ""))
        det.cell(row=n, column=4, value=row.get("attribute", ""))
        det.cell(row=n, column=5, value=row.get("impact", ""))
        det.cell(row=n, column=6, value=to_int(row.get("lines", "")))
        # Band is a formula so editing a threshold in Assumptions re-bands live.
        if basis == "PGM_CHANGE":
            det.cell(row=n, column=7, value=(
                f'=IF($F{n}="","S",'
                f'IF($F{n}<=Assumptions!$B${A_BAND_S},"S",'
                f'IF($F{n}<=Assumptions!$B${A_BAND_M},"M","L")))'
            ))
            det.cell(row=n, column=9, value=(
                f'=IF($G{n}="S",Assumptions!$B${A_PGM_S},'
                f'IF($G{n}="M",Assumptions!$B${A_PGM_M},Assumptions!$B${A_PGM_L}))'
            ))
        elif basis == "NO_COST":
            det.cell(row=n, column=9, value=0)
        else:
            det.cell(row=n, column=9,
                     value=f"=Assumptions!$B${RATE_ROW[basis]}")
        det.cell(row=n, column=8, value=basis)
        det.cell(row=n, column=9).number_format = "0.00"
        det.cell(row=n, column=10, value=row.get("notes", ""))

    last = len(rows) + 1
    det.freeze_panes = "A2"
    if rows:
        det.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    for c, name in enumerate(headers, start=1):
        widest = max([len(name)] + [
            len(str(det.cell(row=r, column=c).value or "")) for r in range(2, last + 1)
        ]) if rows else len(name)
        det.column_dimensions[get_column_letter(c)].width = min(max(widest + 2, 9), 55)

    span = f"'Estimate Detail'!$H$2:$H${max(last, 2)}"
    hours_span = f"'Estimate Detail'!$I$2:$I${max(last, 2)}"

    # ----------------------------------------------------------- Assumptions
    asm = wb.create_sheet("Assumptions")
    asm.append(["Assumption", "Value", "Notes"])
    for c in range(1, 4):
        asm.cell(row=1, column=c).font = head_font
        asm.cell(row=1, column=c).fill = head_fill
    for label, value, note in ASSUMPTIONS:
        asm.append([label, value, note])
    for r in range(2, len(ASSUMPTIONS) + 2):
        cell = asm.cell(row=r, column=2)
        if cell.value is None:
            continue
        cell.fill = money_fill
        cell.font = bold
        cell.number_format = "0%" if r in (A_SYSTEST, A_UAT, A_CONTINGENCY) else "0.00"
    for r in (A_BAND_S, A_BAND_M, A_HOURS_DAY):
        asm.cell(row=r, column=2).number_format = "0"
    asm.column_dimensions["A"].width = 34
    asm.column_dimensions["B"].width = 12
    asm.column_dimensions["C"].width = 52
    note_at = len(ASSUMPTIONS) + 3
    asm.cell(row=note_at, column=1,
             value="Every shaded value feeds the Detail and Summary sheets by "
                   "cell reference. Change one and the totals recalculate — "
                   "there is no need to regenerate the workbook.").font = dim
    asm.merge_cells(start_row=note_at, start_column=1, end_row=note_at, end_column=3)
    asm.cell(row=note_at, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    asm.row_dimensions[note_at].height = 30

    # --------------------------------------------------------------- Summary
    s = wb.create_sheet("Summary")
    s["A1"] = title
    s["A1"].font = Font(bold=True, size=14)
    s["A2"] = "Change Impact Estimate"
    s["A2"].font = dim
    s["A4"], s["B4"] = "Generated", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    s["A5"], s["B5"] = "Author", "iA by programmers.io"
    s["A6"], s["B6"] = "Total impacted objects", len(rows)
    # No per-basis counts here — the Basis table at the foot of this sheet
    # already breaks the total down, and two places to read the same number is
    # two places to disagree.

    r = 8
    s.cell(row=r, column=1, value="Phase").font = head_font
    s.cell(row=r, column=1).fill = head_fill
    s.cell(row=r, column=2, value="Hours").font = head_font
    s.cell(row=r, column=2).fill = head_fill

    design, dev, unit, conv = r + 1, r + 2, r + 3, r + 4
    build_sub, systest, uat, cutover, sub, cont, total, days = (
        r + 5, r + 6, r + 7, r + 8, r + 9, r + 10, r + 11, r + 12
    )
    plan = [
        (design, "Analysis & change design", f"=Assumptions!$B${A_DESIGN}"),
        (dev, "Development", f"=SUM({hours_span})"),
        (unit, "Unit test",
         f'=COUNTIF({span},"PGM_CHANGE")*Assumptions!$B${A_UNIT}'),
        (conv, "Data conversion",
         f"=Assumptions!$B${A_DATACONV}" if narrowing else "0"),
        (build_sub, "Build subtotal", f"=SUM(B{design}:B{conv})"),
        (systest, "System / regression test", f"=B{dev}*Assumptions!$B${A_SYSTEST}"),
        (uat, "UAT support", f"=B{dev}*Assumptions!$B${A_UAT}"),
        (cutover, "Cutover / implementation", f"=Assumptions!$B${A_CUTOVER}"),
        (sub, "Subtotal", f"=B{build_sub}+SUM(B{systest}:B{cutover})"),
        (cont, "Contingency", f"=B{sub}*Assumptions!$B${A_CONTINGENCY}"),
        (total, "TOTAL HOURS", f"=B{sub}+B{cont}"),
        (days, "TOTAL DAYS", f"=B{total}/Assumptions!$B${A_HOURS_DAY}"),
    ]
    for at, label, formula in plan:
        s.cell(row=at, column=1, value=label)
        s.cell(row=at, column=2, value=formula).number_format = "0.00"
        if at in (build_sub, sub, total, days):
            s.cell(row=at, column=1).font = bold
            s.cell(row=at, column=2).font = bold
    s.cell(row=total, column=2).fill = money_fill
    s.cell(row=days, column=2).fill = money_fill

    r = days + 2
    s.cell(row=r, column=1, value="Basis").font = head_font
    s.cell(row=r, column=1).fill = head_fill
    s.cell(row=r, column=2, value="Count").font = head_font
    s.cell(row=r, column=2).fill = head_fill
    s.cell(row=r, column=3, value="Hours").font = head_font
    s.cell(row=r, column=3).fill = head_fill
    for basis in sorted(counts):
        r += 1
        s.cell(row=r, column=1, value=basis)
        s.cell(row=r, column=2, value=f'=COUNTIF({span},"{basis}")')
        s.cell(row=r, column=3,
               value=f'=SUMIF({span},"{basis}",{hours_span})').number_format = "0.00"

    r += 2
    s.cell(row=r, column=1, value=(
        "Rates are on the Assumptions sheet and are deliberately lean. Every "
        "figure here is a formula — edit a rate and this page updates. Excludes "
        "data validation against live data, external interfaces, and anything "
        "changed since the last iA metadata build."
    )).font = dim
    s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    s.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    s.row_dimensions[r].height = 44
    s.column_dimensions["A"].width = 32
    s.column_dimensions["B"].width = 14
    s.column_dimensions["C"].width = 14
    s.column_dimensions["D"].width = 14

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    # Assumptions leads: the rates are what a client argues about, so the
    # workbook opens on the negotiable page rather than on the total.
    wb._sheets = [asm, s, det]
    wb.active = 0
    wb.save(out)
    return counts


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("markdown", help="the Change Impact Assessment .md")
    ap.add_argument("--xlsx", action="store_true",
                    help="write the estimate workbook (default action)")
    ap.add_argument("--narrowing", action="store_true",
                    help="the change can lose data — adds the data-conversion line")
    ap.add_argument("--outdir", help="output directory (default: beside the .md)")
    args = ap.parse_args()

    src = Path(args.markdown)
    if not src.is_file():
        print(f"error: not found: {src}", file=sys.stderr)
        return 1

    title, rows = parse(src)
    if not rows:
        print(f"error: no '## Affected Artifacts' table found in {src.name} — "
              "see references/change-impact-analysis.md §7 for the required "
              "columns", file=sys.stderr)
        return 1

    outdir = Path(args.outdir) if args.outdir else src.parent
    outdir.mkdir(parents=True, exist_ok=True)
    out = outdir / f"{src.stem.replace('_Change_Impact', '')}_Change_Estimate.xlsx"

    counts = build(out, title, rows, args.narrowing)
    print(f"{src.name}: {len(rows)} artifacts")
    for basis in sorted(counts):
        print(f"  {basis:<18} {counts[basis]}")
    print(f"  xlsx -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
