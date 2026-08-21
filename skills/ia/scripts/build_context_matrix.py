#!/usr/bin/env python3
"""Render an Object Context Matrix markdown report into Excel and/or HTML.

The markdown written by the /ia skill is the source of truth. This script only
reformats it, so whatever the agent chose to include — extra sections, a
reordered summary, enrichment tables — lands in every download unchanged. It
never queries the repository, which is exactly what keeps the file the user
downloaded identical to the answer they saw on screen.

Input shape (produced by references/object-context-matrix.md):

    # Object Context Matrix — CUSTMNTR (*PGM, CASELIB)
    <summary paragraphs>
    ## Files Used In Update Mode(s) (1)
    | File Library | File Name | ... |
    |---|---|---|
    | CASELIB | CUSTMST | ... |
    ## Display File(s) (1)
    ...

Outputs:
  --xlsx   one sheet per ## section, frozen header, autofilter, sized columns;
           prose (Summary, Notes, per-section commentary) is carried across too
  --html   one self-contained dark-themed page: collapsible sections and a live
           filter box, mirroring the portal screen. No external assets.

Requires:  pip install openpyxl   (only for --xlsx; --html has no dependencies)
Usage:     python build_context_matrix.py REPORT.md [--xlsx] [--html]
                                          [--outdir DIR]
"""
from __future__ import annotations

import argparse
import datetime as _dt
import html
import re
import sys
from pathlib import Path

MAX_SHEET_NAME = 31
# Excel forbids these in a sheet name.
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")
# Headings carry their row count as "Display File(s) — 1"; older reports used
# "Display File(s) (1)". Either form gets dropped from the sheet tab.
_TRAILING_COUNT = re.compile(r"\s*(?:\(\d+\)|[-–—]\s*\d+)\s*$")
# A markdown bullet or numbered item — anything else continues the line above.
_BULLET = re.compile(r"^(?:[-*+]|\d+[.)])\s")


def split_row(line: str) -> list[str]:
    """Split a markdown table row, honouring escaped pipes."""
    line = line.strip().replace(r"\|", "\x00")
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.replace("\x00", "|").strip() for c in line.split("|")]


def is_divider(line: str) -> bool:
    """True for the |---|---| row under a table header."""
    cells = split_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "")


# IBM i object types (*PGM, *FILE, *MENU, *ALL, *ENTRY…) start with an asterisk,
# which the emphasis stripper below would otherwise eat — turning *PGM into PGM
# in every Object Type column. Park them before stripping, restore after.
_OBJ_TYPE = re.compile(r"\*([A-Z][A-Z0-9]{1,9})\b")
_CODE = re.compile(r"`([^`]*)`")
# Underscores only mark emphasis at a word boundary. Without this guard the
# stripper turns detected_from into detectedfrom and ia_file_dependencies into
# iafiledependencies — silently corrupting every identifier it touches.
_EMPH = re.compile(r"\*\*|\*|(?<!\w)__?(?=\S)|(?<=\S)__?(?!\w)")


def strip_md(text: str) -> str:
    """Flatten inline markdown — cells land in Excel as plain text."""
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    # Park inline code so its contents pass through untouched.
    spans: list[str] = []

    def hold(m: re.Match) -> str:
        spans.append(m.group(1))
        return f"\x02{len(spans) - 1}\x02"

    text = _CODE.sub(hold, text)
    text = _OBJ_TYPE.sub("\x01\\1", text)
    text = _EMPH.sub("", text)
    text = text.replace("\x01", "*")
    text = re.sub(r"\x02(\d+)\x02", lambda m: spans[int(m.group(1))], text)
    return text.strip()


def parse(md_path: Path) -> tuple[str, list[str], list[dict]]:
    """Return (title, summary_paragraphs, sections)."""
    lines = md_path.read_text(encoding="utf-8").splitlines()

    title = md_path.stem.replace("_", " ")
    summary: list[str] = []
    sections: list[dict] = []
    current: dict | None = None
    # Markdown hard-wraps prose, so a paragraph arrives as several lines. Only a
    # blank line (or a heading/table) ends one — without this, every wrapped
    # bullet lands as its own half-sentence row.
    fresh = True
    i = 0

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if line.startswith("# ") and not line.startswith("##"):
            title = strip_md(line[2:])
            fresh = True
            i += 1
            continue

        if line.startswith("## "):
            current = {"heading": strip_md(line[3:]), "header": [], "rows": [], "notes": []}
            sections.append(current)
            fresh = True
            i += 1
            continue

        # A table starts at a pipe row followed by a divider row.
        if line.startswith("|") and i + 1 < len(lines) and is_divider(lines[i + 1]):
            header = [strip_md(c) for c in split_row(line)]
            rows: list[list[str]] = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_divider(lines[i]):
                    rows.append([strip_md(c) for c in split_row(lines[i])])
                i += 1
            # Pad or trim ragged rows to the header width.
            width = len(header)
            rows = [(r + [""] * width)[:width] for r in rows]
            if current is None:
                current = {"heading": "Details", "header": [], "rows": [], "notes": []}
                sections.append(current)
            if current["header"]:
                # A second table under one heading — keep it as its own section.
                current = {"heading": current["heading"] + " (cont.)",
                           "header": [], "rows": [], "notes": []}
                sections.append(current)
            current["header"] = header
            current["rows"] = rows
            fresh = True
            continue

        if line:
            target = summary if current is None else current["notes"]
            text = strip_md(line)
            if not fresh and target and not _BULLET.match(line):
                target[-1] = f"{target[-1]} {text}"
            else:
                target.append(text)
            fresh = False
        else:
            fresh = True
        i += 1

    return title, summary, [s for s in sections if s["rows"] or s["notes"]]


def clip(text: str, cap: int) -> str:
    """Trim to cap on a word boundary rather than slicing a word in half."""
    if len(text) <= cap:
        return text
    cut = text[:cap]
    space = cut.rfind(" ")
    # Only honour the boundary if it keeps most of the name; otherwise a single
    # long word would collapse the tab to nothing.
    if space >= cap * 0.6:
        cut = cut[:space]
    return cut.rstrip(" -–—([")


def sheet_name(heading: str, used: set[str]) -> str:
    # Drop the trailing row count — Overview lists it anyway, and keeping it
    # pushes longer headings past Excel's 31-char cap.
    name = _TRAILING_COUNT.sub("", heading)
    name = _SHEET_BAD.sub("-", name).strip() or "Section"
    name = clip(name, MAX_SHEET_NAME)
    base, n = name, 2
    while name.lower() in used:
        suffix = f" {n}"
        name = clip(base, MAX_SHEET_NAME - len(suffix)) + suffix
        n += 1
    used.add(name.lower())
    return name


def write_xlsx(out: Path, title: str, summary: list[str], sections: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("openpyxl is required for --xlsx:  pip install openpyxl")

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Generated"
    ws["B3"] = _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    ws["A4"] = "Author"
    ws["B4"] = "iA by programmers.io"
    row = 6
    for para in summary:
        ws.cell(row=row, column=1, value=para).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Section").font = head_font
    ws.cell(row=row, column=1).fill = head_fill
    ws.cell(row=row, column=2, value="Rows").font = head_font
    ws.cell(row=row, column=2).fill = head_fill
    for sec in sections:
        row += 1
        ws.cell(row=row, column=1, value=sec["heading"])
        # Prose-only sections (Summary, Notes) have no rows to count — leaving
        # the cell blank reads better than a misleading 0.
        ws.cell(row=row, column=2, value=len(sec["rows"]) if sec["rows"] else "")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 10

    note_font = Font(italic=True, color="44546A")

    def put_notes(sheet, start: int, notes: list[str], span: int) -> None:
        """Lay prose out down column A, merged across the table's width."""
        for n, text in enumerate(notes):
            at = start + n
            cell = sheet.cell(row=at, column=1, value=text)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = note_font
            if span > 1:
                sheet.merge_cells(start_row=at, start_column=1,
                                  end_row=at, end_column=span)
            sheet.row_dimensions[at].height = 15 * max(1, (len(text) // (span * 18) + 1))

    used: set[str] = {"overview"}
    for sec in sections:
        if not (sec["rows"] or sec["notes"]):
            continue
        s = wb.create_sheet(sheet_name(sec["heading"], used))

        if not sec["rows"]:
            # Prose-only section — the heading is the sheet's only chrome.
            s.cell(row=1, column=1, value=sec["heading"]).font = Font(bold=True, size=12)
            s.column_dimensions["A"].width = 110
            put_notes(s, 3, sec["notes"], 1)
            continue

        s.append(sec["header"])
        for c in range(1, len(sec["header"]) + 1):
            s.cell(row=1, column=c).font = head_font
            s.cell(row=1, column=c).fill = head_fill
        for r in sec["rows"]:
            s.append(r)
        s.freeze_panes = "A2"
        s.auto_filter.ref = (
            f"A1:{get_column_letter(len(sec['header']))}{len(sec['rows']) + 1}"
        )
        for c, name in enumerate(sec["header"], start=1):
            longest = max([len(name)] + [len(str(r[c - 1])) for r in sec["rows"]])
            s.column_dimensions[get_column_letter(c)].width = min(max(longest + 2, 10), 60)
        # A blank row keeps the prose clear of the autofilter range above it.
        put_notes(s, len(sec["rows"]) + 3, sec["notes"], len(sec["header"]))

    wb.save(out)


HTML_CSS = """
:root{--bg:#0f1419;--panel:#161d26;--edge:#263140;--fg:#e6edf3;--dim:#8b98a5;
--accent:#4a9eff;--head:#1c2734}
*{box-sizing:border-box}
body{margin:0;padding:2rem 1.25rem 4rem;background:var(--bg);color:var(--fg);
font:15px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif}
.wrap{max-width:1200px;margin:0 auto}
h1{font-size:1.6rem;margin:0 0 .35rem;font-weight:650}
.sub{color:var(--dim);font-size:.85rem;margin-bottom:1.5rem}
.summary{background:var(--panel);border:1px solid var(--edge);border-left:3px solid
var(--accent);border-radius:6px;padding:.9rem 1.1rem;margin-bottom:1.5rem}
.summary p{margin:.4rem 0}
#filter{width:100%;padding:.6rem .85rem;margin-bottom:1.25rem;background:var(--panel);
border:1px solid var(--edge);border-radius:6px;color:var(--fg);font-size:.9rem}
#filter:focus{outline:none;border-color:var(--accent)}
details{background:var(--panel);border:1px solid var(--edge);border-radius:6px;
margin-bottom:.75rem;overflow:hidden}
summary{padding:.75rem 1.1rem;cursor:pointer;font-weight:600;background:var(--head);
display:flex;align-items:center;gap:.6rem;list-style:none}
summary::-webkit-details-marker{display:none}
summary::before{content:"\\25B8";color:var(--accent);transition:transform .15s}
details[open] summary::before{transform:rotate(90deg)}
.count{background:var(--accent);color:#0f1419;border-radius:10px;padding:.05rem .5rem;
font-size:.75rem;font-weight:700}
.scroll{overflow-x:auto}
table{border-collapse:collapse;width:100%;font-size:.86rem}
th{background:var(--head);text-align:left;padding:.55rem .8rem;white-space:nowrap;
border-bottom:1px solid var(--edge);position:sticky;top:0}
td{padding:.5rem .8rem;border-bottom:1px solid var(--edge);vertical-align:top}
tr:last-child td{border-bottom:none}
tr:hover td{background:#1b2430}
.note{color:var(--dim);padding:.6rem 1.1rem;font-size:.85rem}
footer{margin-top:2.5rem;color:var(--dim);font-size:.8rem;text-align:center}
.hidden{display:none}
@media print{body{background:#fff;color:#000}details{break-inside:avoid}
details>div{display:block!important}#filter{display:none}}
"""

HTML_JS = """
var box=document.getElementById('filter');
box.addEventListener('input',function(){
  var q=box.value.toLowerCase();
  document.querySelectorAll('details').forEach(function(d){
    var hits=0;
    d.querySelectorAll('tbody tr').forEach(function(tr){
      var on=!q||tr.textContent.toLowerCase().indexOf(q)>-1;
      tr.classList.toggle('hidden',!on);
      if(on)hits++;
    });
    var tot=d.querySelectorAll('tbody tr').length;
    d.classList.toggle('hidden',q&&tot&&!hits);
    if(q&&hits)d.open=true;
  });
});
"""


def write_html(out: Path, title: str, summary: list[str], sections: list[dict]) -> None:
    e = html.escape
    p: list[str] = []
    p.append("<!doctype html><html lang='en'><head><meta charset='utf-8'>")
    p.append("<meta name='viewport' content='width=device-width,initial-scale=1'>")
    p.append(f"<title>{e(title)}</title><style>{HTML_CSS}</style></head><body><div class='wrap'>")
    p.append(f"<h1>{e(title)}</h1>")
    p.append("<div class='sub'>Generated "
             f"{_dt.datetime.now().strftime('%Y-%m-%d %H:%M')} &middot; "
             "iA by programmers.io</div>")
    if summary:
        p.append("<div class='summary'>")
        p.extend(f"<p>{e(s)}</p>" for s in summary)
        p.append("</div>")
    p.append("<input id='filter' type='search' placeholder='Filter every section…'>")

    for sec in sections:
        p.append("<details open><summary>" + e(sec["heading"]))
        if sec["rows"]:
            p.append(f"<span class='count'>{len(sec['rows'])}</span>")
        p.append("</summary><div>")
        for note in sec["notes"]:
            p.append(f"<div class='note'>{e(note)}</div>")
        if sec["rows"]:
            p.append("<div class='scroll'><table><thead><tr>")
            p.extend(f"<th>{e(h)}</th>" for h in sec["header"])
            p.append("</tr></thead><tbody>")
            for r in sec["rows"]:
                p.append("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>")
            p.append("</tbody></table></div>")
        p.append("</div></details>")

    p.append("<footer>Object Context Matrix &middot; iA by programmers.io</footer>")
    p.append(f"</div><script>{HTML_JS}</script></body></html>")
    out.write_text("".join(p), encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("markdown", help="the Object Context Matrix .md report")
    ap.add_argument("--xlsx", action="store_true", help="write the Excel workbook")
    ap.add_argument("--html", action="store_true", help="write the interactive page")
    ap.add_argument("--outdir", help="output directory (default: beside the .md)")
    args = ap.parse_args()

    src = Path(args.markdown)
    if not src.is_file():
        return err(f"not found: {src}")
    if not (args.xlsx or args.html):
        return err("nothing to do — pass --xlsx and/or --html")

    title, summary, sections = parse(src)
    if not sections:
        return err(f"no '## ' sections with tables found in {src.name} — "
                   "is this an Object Context Matrix report?")

    outdir = Path(args.outdir) if args.outdir else src.parent
    outdir.mkdir(parents=True, exist_ok=True)

    total = sum(len(s["rows"]) for s in sections)
    print(f"{src.name}: {len(sections)} sections, {total} rows")

    if args.xlsx:
        out = outdir / f"{src.stem}.xlsx"
        write_xlsx(out, title, summary, sections)
        print(f"  xlsx -> {out}")
    if args.html:
        out = outdir / f"{src.stem}.html"
        write_html(out, title, summary, sections)
        print(f"  html -> {out}")
    return 0


def err(msg: str) -> int:
    print(f"error: {msg}", file=sys.stderr)
    return 1


if __name__ == "__main__":
    sys.exit(main())
